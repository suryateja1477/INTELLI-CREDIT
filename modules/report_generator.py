"""
=============================================================
MODULE: report_generator.py
PURPOSE: Generates a professional Microsoft Word (.docx)
         Credit Appraisal Memo (CAM) — the final deliverable
         that a bank credit officer would present to the
         loan sanctioning committee.

REPORT SECTIONS (9 total — mirrors real Indian bank CAMs):
  1. Company Overview
  2. Character Assessment
  3. Capacity Assessment
  4. Capital Assessment
  5. Collateral Assessment
  6. Conditions Assessment
  7. GST & Fraud Risk Analysis
  8. Early Warning Signals
  9. FINAL RECOMMENDATION

OUTPUT: Saved as a .docx file in the outputs/ folder.
=============================================================
"""

import os
import json
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def generate_cam_report(credit_decision, financial_data, gst_analysis, research_summary, officer_assessment=None, company_name="Company"):
    """
    MAIN FUNCTION — Call this from app.py

    Generates a professional Word document (.docx) containing the
    complete Credit Appraisal Memo (CAM) with all 9 sections.

    Args:
        credit_decision:   Dict from scoring_engine.py — final scores & decision
        financial_data:    Dict from pdf_reader.py — extracted financial data
        gst_analysis:      Dict from gst_checker.py — GST fraud analysis
        research_summary:  Dict from research_agent.py — web research findings
        officer_assessment: Dict from research_agent.py — officer notes (optional)
        company_name:      Name of the company being appraised

    Returns:
        The file path of the saved Word document (string).
        Returns a dict with "error" key if something goes wrong.

    Example usage in app.py:
        file_path = generate_cam_report(
            credit_decision, financial_data,
            gst_analysis, research_summary,
            officer_assessment, "Sharma Textiles Pvt Ltd"
        )
        with open(file_path, "rb") as f:
            st.download_button("Download CAM Report", f, file_name="CAM_Report.docx")
    """
    try:
        # Create a new Word document
        doc = Document()

        # Set default font for the whole document
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Calibri"
        font.size = Pt(11)
        font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        # ============================
        # COVER PAGE / HEADER
        # ============================
        add_cover_section(doc, company_name, credit_decision)

        # ============================
        # SECTION 1: Company Overview
        # ============================
        add_section_1_company_overview(doc, financial_data, company_name)

        # ============================
        # SECTION 2: Character Assessment
        # ============================
        add_section_2_character(doc, credit_decision, research_summary)

        # ============================
        # SECTION 3: Capacity Assessment
        # ============================
        add_section_3_capacity(doc, credit_decision, financial_data)

        # ============================
        # SECTION 4: Capital Assessment
        # ============================
        add_section_4_capital(doc, credit_decision, financial_data)

        # ============================
        # SECTION 5: Collateral Assessment
        # ============================
        add_section_5_collateral(doc, credit_decision)

        # ============================
        # SECTION 6: Conditions Assessment
        # ============================
        add_section_6_conditions(doc, credit_decision)

        # ============================
        # SECTION 7: GST & Fraud Risk Analysis
        # ============================
        add_section_7_gst_fraud(doc, gst_analysis)

        # ============================
        # SECTION 8: Early Warning Signals
        # ============================
        add_section_8_early_warnings(doc, credit_decision, gst_analysis, research_summary)

        # ============================
        # SECTION 9: FINAL RECOMMENDATION
        # ============================
        add_section_9_final_recommendation(doc, credit_decision, officer_assessment)

        # ============================
        # FOOTER — Disclaimer
        # ============================
        add_disclaimer(doc)

        # Create outputs folder if it doesn't exist
        os.makedirs("outputs", exist_ok=True)

        # Generate filename with timestamp
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_name = company_name.replace(" ", "_").replace("/", "_")[:30]
        filename = f"CAM_Report_{safe_name}_{timestamp}.docx"
        filepath = os.path.join("outputs", filename)

        # Save the document
        doc.save(filepath)

        return filepath

    except Exception as e:
        return {"error": f"Report generation failed: {str(e)}"}


# =============================================================
# HELPER: Cover Section
# =============================================================
def add_cover_section(doc, company_name, credit_decision):
    """Adds the title page / header section of the CAM report."""
    try:
        # Bank header
        header_para = doc.add_paragraph()
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = header_para.add_run("CREDIT APPRAISAL MEMORANDUM")
        run.bold = True
        run.font.size = Pt(22)
        run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

        # Subtitle
        sub_para = doc.add_paragraph()
        sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = sub_para.add_run("Confidential — For Internal Use Only")
        run.italic = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

        # Divider line
        doc.add_paragraph("━" * 60)

        # Company name
        name_para = doc.add_paragraph()
        name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = name_para.add_run(company_name.upper())
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

        # Date and reference
        info_para = doc.add_paragraph()
        info_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_str = datetime.now().strftime("%d %B %Y")
        run = info_para.add_run(f"Date of Appraisal: {date_str}")
        run.font.size = Pt(12)

        # Decision banner
        decision = safe_get(credit_decision, "decision", "PENDING")
        overall_score = safe_get(credit_decision, "overall_score", "N/A")

        banner_para = doc.add_paragraph()
        banner_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = banner_para.add_run(f"\n★  DECISION: {decision}  |  SCORE: {overall_score}/100  ★\n")
        run.bold = True
        run.font.size = Pt(16)

        # Color the decision text
        if "APPROVE" == str(decision).upper():
            run.font.color.rgb = RGBColor(0x27, 0xAE, 0x60)  # Green
        elif "REJECT" == str(decision).upper():
            run.font.color.rgb = RGBColor(0xE7, 0x4C, 0x3C)  # Red
        else:
            run.font.color.rgb = RGBColor(0xF3, 0x9C, 0x12)  # Orange

        # Another divider
        doc.add_paragraph("━" * 60)

        # Summary table — quick overview
        summary_table = doc.add_table(rows=4, cols=2)
        summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        summary_table.style = "Light Grid Accent 1"

        summary_data = [
            ("Recommended Loan Amount", safe_get(credit_decision, "recommended_loan_amount", "N/A")),
            ("Recommended Interest Rate", safe_get(credit_decision, "recommended_interest_rate", "N/A")),
            ("Recommended Tenure", safe_get(credit_decision, "recommended_tenure", "N/A")),
            ("Decision Confidence", safe_get(credit_decision, "decision_confidence", "N/A"))
        ]

        for i, (label, value) in enumerate(summary_data):
            summary_table.rows[i].cells[0].text = label
            summary_table.rows[i].cells[1].text = str(value)
            # Bold the label cells
            for paragraph in summary_table.rows[i].cells[0].paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        doc.add_paragraph("")  # spacing

    except Exception as e:
        doc.add_paragraph(f"[Cover section error: {str(e)}]")


# =============================================================
# SECTION 1: Company Overview
# =============================================================
def add_section_1_company_overview(doc, financial_data, company_name):
    """Section 1 — Company Overview with key financial snapshot."""
    try:
        add_section_heading(doc, "1. COMPANY OVERVIEW")

        # Company name
        doc.add_paragraph(f"Company Name: {company_name}", style="List Bullet")

        # Extract key data from financial_data
        fy = safe_get(financial_data, "financial_year", "Not Available")
        doc_type = safe_get(financial_data, "document_type", "Not Available")
        doc.add_paragraph(f"Financial Year: {fy}", style="List Bullet")
        doc.add_paragraph(f"Document Type Analyzed: {doc_type}", style="List Bullet")

        # Financial snapshot table
        doc.add_paragraph("")
        add_sub_heading(doc, "Financial Snapshot")

        fin_table = doc.add_table(rows=7, cols=2)
        fin_table.style = "Light Grid Accent 1"

        # Revenue
        revenue = safe_get(financial_data, "revenue", {})
        rev_val = safe_get(revenue, "value", "N/A") if isinstance(revenue, dict) else str(revenue)
        rev_unit = safe_get(revenue, "unit", "") if isinstance(revenue, dict) else ""

        # Net Profit
        profit = safe_get(financial_data, "net_profit", {})
        profit_val = safe_get(profit, "value", "N/A") if isinstance(profit, dict) else str(profit)
        profit_unit = safe_get(profit, "unit", "") if isinstance(profit, dict) else ""

        # Total Debt
        debt = safe_get(financial_data, "total_debt", {})
        debt_val = safe_get(debt, "value", "N/A") if isinstance(debt, dict) else str(debt)
        debt_unit = safe_get(debt, "unit", "") if isinstance(debt, dict) else ""

        # Net Worth
        nw = safe_get(financial_data, "net_worth", {})
        nw_val = safe_get(nw, "value", "N/A") if isinstance(nw, dict) else str(nw)
        nw_unit = safe_get(nw, "unit", "") if isinstance(nw, dict) else ""

        # Total Assets
        assets = safe_get(financial_data, "total_assets", {})
        assets_val = safe_get(assets, "value", "N/A") if isinstance(assets, dict) else str(assets)
        assets_unit = safe_get(assets, "unit", "") if isinstance(assets, dict) else ""

        # Key Ratios
        ratios = safe_get(financial_data, "key_ratios", {})
        de_ratio = safe_get(ratios, "debt_to_equity", "N/A") if isinstance(ratios, dict) else "N/A"
        current_ratio = safe_get(ratios, "current_ratio", "N/A") if isinstance(ratios, dict) else "N/A"

        fin_data = [
            ("Revenue / Turnover", f"{rev_val} {rev_unit}".strip()),
            ("Net Profit (PAT)", f"{profit_val} {profit_unit}".strip()),
            ("Total Debt", f"{debt_val} {debt_unit}".strip()),
            ("Net Worth", f"{nw_val} {nw_unit}".strip()),
            ("Total Assets", f"{assets_val} {assets_unit}".strip()),
            ("Debt-to-Equity Ratio", str(de_ratio)),
            ("Current Ratio", str(current_ratio))
        ]

        for i, (label, value) in enumerate(fin_data):
            fin_table.rows[i].cells[0].text = label
            fin_table.rows[i].cells[1].text = str(value)
            for paragraph in fin_table.rows[i].cells[0].paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        # Auditor Details
        auditor = safe_get(financial_data, "auditor_details", {})
        if isinstance(auditor, dict) and auditor:
            doc.add_paragraph("")
            add_sub_heading(doc, "Auditor Details")
            doc.add_paragraph(f"Auditor: {safe_get(auditor, 'auditor_name', 'N/A')}", style="List Bullet")
            doc.add_paragraph(f"Opinion: {safe_get(auditor, 'audit_opinion', 'N/A')}", style="List Bullet")
            changed = safe_get(auditor, "auditor_changed", "N/A")
            doc.add_paragraph(f"Auditor Changed Recently: {changed}", style="List Bullet")
            qualifications = safe_get(auditor, "key_qualifications", "N/A")
            if qualifications and qualifications != "N/A":
                doc.add_paragraph(f"Key Qualifications: {qualifications}", style="List Bullet")

        doc.add_paragraph("")

    except Exception as e:
        doc.add_paragraph(f"[Section 1 error: {str(e)}]")


# =============================================================
# SECTION 2: Character Assessment
# =============================================================
def add_section_2_character(doc, credit_decision, research_summary):
    """Section 2 — Character Assessment (promoter trust, governance)."""
    try:
        add_section_heading(doc, "2. CHARACTER ASSESSMENT")

        char_data = safe_get(credit_decision, "character_assessment", {})
        add_score_box(doc, "Character Score", char_data)

        # Promoter research findings
        if isinstance(research_summary, dict) and "error" not in research_summary:
            promoter = safe_get(research_summary, "promoter_assessment", {})
            if isinstance(promoter, dict) and promoter:
                add_sub_heading(doc, "Promoter Background (from External Research)")
                rep_score = safe_get(promoter, "reputation_score", "N/A")
                doc.add_paragraph(f"Promoter Reputation Score: {rep_score}/10", style="List Bullet")

                controversies = safe_get(promoter, "controversies_found", False)
                doc.add_paragraph(f"Controversies Found: {'Yes' if controversies else 'No'}", style="List Bullet")

                if controversies:
                    details = safe_get(promoter, "controversy_details", "")
                    if details:
                        doc.add_paragraph(f"Details: {details}", style="List Bullet 2")

                linked = safe_get(promoter, "linked_to_other_defaulting_companies", False)
                doc.add_paragraph(f"Linked to Defaulting Companies: {'Yes' if linked else 'No'}", style="List Bullet")

        doc.add_paragraph("")

    except Exception as e:
        doc.add_paragraph(f"[Section 2 error: {str(e)}]")


# =============================================================
# SECTION 3: Capacity Assessment
# =============================================================
def add_section_3_capacity(doc, credit_decision, financial_data):
    """Section 3 — Capacity Assessment (ability to repay)."""
    try:
        add_section_heading(doc, "3. CAPACITY ASSESSMENT")

        cap_data = safe_get(credit_decision, "capacity_assessment", {})
        add_score_box(doc, "Capacity Score", cap_data)

        # Cash flow details
        cash_flow = safe_get(financial_data, "cash_flow_from_operations", {})
        if isinstance(cash_flow, dict) and cash_flow:
            add_sub_heading(doc, "Cash Flow Analysis")
            cf_val = safe_get(cash_flow, "value", "N/A")
            cf_positive = safe_get(cash_flow, "positive", "N/A")
            doc.add_paragraph(f"Operating Cash Flow: {cf_val}", style="List Bullet")
            doc.add_paragraph(f"Cash Flow Positive: {cf_positive}", style="List Bullet")

        # DSCR
        dscr = safe_get(cap_data, "dscr_assessment", "")
        if dscr:
            doc.add_paragraph(f"DSCR Assessment: {dscr}", style="List Bullet")

        doc.add_paragraph("")

    except Exception as e:
        doc.add_paragraph(f"[Section 3 error: {str(e)}]")


# =============================================================
# SECTION 4: Capital Assessment
# =============================================================
def add_section_4_capital(doc, credit_decision, financial_data):
    """Section 4 — Capital Assessment (net worth, leverage)."""
    try:
        add_section_heading(doc, "4. CAPITAL ASSESSMENT")

        capital_data = safe_get(credit_decision, "capital_assessment", {})
        add_score_box(doc, "Capital Score", capital_data)

        # Debt-to-equity from decision
        de = safe_get(capital_data, "debt_to_equity", "")
        if de:
            doc.add_paragraph(f"Debt-to-Equity Assessment: {de}", style="List Bullet")

        nw_adequacy = safe_get(capital_data, "net_worth_adequacy", "")
        if nw_adequacy:
            doc.add_paragraph(f"Net Worth Adequacy: {nw_adequacy}", style="List Bullet")

        doc.add_paragraph("")

    except Exception as e:
        doc.add_paragraph(f"[Section 4 error: {str(e)}]")


# =============================================================
# SECTION 5: Collateral Assessment
# =============================================================
def add_section_5_collateral(doc, credit_decision):
    """Section 5 — Collateral Assessment (assets pledged)."""
    try:
        add_section_heading(doc, "5. COLLATERAL ASSESSMENT")

        coll_data = safe_get(credit_decision, "collateral_assessment", {})
        add_score_box(doc, "Collateral Score", coll_data)

        cover = safe_get(coll_data, "collateral_cover_ratio", "")
        if cover:
            doc.add_paragraph(f"Collateral Cover Ratio: {cover}", style="List Bullet")

        quality = safe_get(coll_data, "asset_quality", "")
        if quality:
            doc.add_paragraph(f"Asset Quality: {quality}", style="List Bullet")

        doc.add_paragraph("")

    except Exception as e:
        doc.add_paragraph(f"[Section 5 error: {str(e)}]")


# =============================================================
# SECTION 6: Conditions Assessment
# =============================================================
def add_section_6_conditions(doc, credit_decision):
    """Section 6 — Conditions Assessment (industry, macro)."""
    try:
        add_section_heading(doc, "6. CONDITIONS ASSESSMENT")

        cond_data = safe_get(credit_decision, "conditions_assessment", {})
        add_score_box(doc, "Conditions Score", cond_data)

        outlook = safe_get(cond_data, "industry_outlook", "")
        if outlook:
            doc.add_paragraph(f"Industry Outlook: {outlook}", style="List Bullet")

        reg_risk = safe_get(cond_data, "regulatory_risk", "")
        if reg_risk:
            doc.add_paragraph(f"Regulatory Risk: {reg_risk}", style="List Bullet")

        doc.add_paragraph("")

    except Exception as e:
        doc.add_paragraph(f"[Section 6 error: {str(e)}]")


# =============================================================
# SECTION 7: GST & Fraud Risk Analysis
# =============================================================
def add_section_7_gst_fraud(doc, gst_analysis):
    """Section 7 — GST cross-check and fraud detection results."""
    try:
        add_section_heading(doc, "7. GST & FRAUD RISK ANALYSIS")

        if isinstance(gst_analysis, dict) and "error" not in gst_analysis:

            # Mismatch analysis
            mismatch = safe_get(gst_analysis, "mismatch_analysis", {})
            if isinstance(mismatch, dict):
                add_sub_heading(doc, "GSTR-3B vs GSTR-2A Mismatch")
                doc.add_paragraph(f"Mismatch Amount: {safe_get(mismatch, 'turnover_mismatch_amount', 'N/A')}", style="List Bullet")
                doc.add_paragraph(f"Mismatch Percentage: {safe_get(mismatch, 'mismatch_percentage', 'N/A')}", style="List Bullet")
                doc.add_paragraph(f"Severity: {safe_get(mismatch, 'severity', 'N/A')}", style="List Bullet")
                interpretation = safe_get(mismatch, "interpretation", "")
                if interpretation:
                    doc.add_paragraph(f"Interpretation: {interpretation}", style="List Bullet 2")

            # Circular trading
            circular = safe_get(gst_analysis, "circular_trading_detection", {})
            if isinstance(circular, dict):
                add_sub_heading(doc, "Circular Trading Detection")
                detected = safe_get(circular, "detected", False)
                doc.add_paragraph(f"Circular Trading Detected: {'YES — ⚠️' if detected else 'No'}", style="List Bullet")
                doc.add_paragraph(f"Confidence: {safe_get(circular, 'confidence', 'N/A')}", style="List Bullet")

                signals = safe_get(circular, "signals_found", [])
                if isinstance(signals, list) and signals:
                    for signal in signals:
                        doc.add_paragraph(f"• {signal}", style="List Bullet 2")

            # Revenue inflation
            inflation = safe_get(gst_analysis, "revenue_inflation_risk", {})
            if isinstance(inflation, dict):
                add_sub_heading(doc, "Revenue Inflation Risk")
                doc.add_paragraph(f"Risk Level: {safe_get(inflation, 'risk_level', 'N/A')}", style="List Bullet")
                explanation = safe_get(inflation, "explanation", "")
                if explanation:
                    doc.add_paragraph(f"Explanation: {explanation}", style="List Bullet 2")

            # Overall GST risk
            overall_gst = safe_get(gst_analysis, "overall_gst_risk", {})
            if isinstance(overall_gst, dict):
                add_sub_heading(doc, "Overall GST Risk Assessment")
                doc.add_paragraph(f"Risk Level: {safe_get(overall_gst, 'risk_level', 'N/A')}", style="List Bullet")
                doc.add_paragraph(f"Risk Score: {safe_get(overall_gst, 'risk_score', 'N/A')}/100", style="List Bullet")

                concerns = safe_get(overall_gst, "top_concerns", [])
                if isinstance(concerns, list) and concerns:
                    doc.add_paragraph("Top Concerns:")
                    for concern in concerns:
                        doc.add_paragraph(f"• {concern}", style="List Bullet 2")

        else:
            doc.add_paragraph("GST analysis data not available or encountered an error.")

        doc.add_paragraph("")

    except Exception as e:
        doc.add_paragraph(f"[Section 7 error: {str(e)}]")


# =============================================================
# SECTION 8: Early Warning Signals
# =============================================================
def add_section_8_early_warnings(doc, credit_decision, gst_analysis, research_summary):
    """Section 8 — Combined early warning signals from all sources."""
    try:
        add_section_heading(doc, "8. EARLY WARNING SIGNALS")

        # Collect all early warning signals from different sources
        all_warnings = []

        # From credit decision
        ews_decision = safe_get(credit_decision, "early_warning_signals", [])
        if isinstance(ews_decision, list):
            all_warnings.extend(ews_decision)

        # From GST analysis
        ews_gst = safe_get(gst_analysis, "early_warning_signals", [])
        if isinstance(ews_gst, list):
            all_warnings.extend(ews_gst)

        # From research
        ews_research = safe_get(research_summary, "early_warning_signals", [])
        if isinstance(ews_research, list):
            all_warnings.extend(ews_research)

        # Remove duplicates while preserving order
        seen = set()
        unique_warnings = []
        for w in all_warnings:
            w_lower = str(w).lower().strip()
            if w_lower not in seen:
                seen.add(w_lower)
                unique_warnings.append(w)

        if unique_warnings:
            # Add warning header
            warning_para = doc.add_paragraph()
            run = warning_para.add_run("⚠️  The following early warning signals have been detected:")
            run.bold = True
            run.font.color.rgb = RGBColor(0xE7, 0x4C, 0x3C)

            doc.add_paragraph("")

            # Create a table for warnings — looks more professional
            warn_table = doc.add_table(rows=len(unique_warnings) + 1, cols=2)
            warn_table.style = "Light Grid Accent 2"

            # Header row
            warn_table.rows[0].cells[0].text = "#"
            warn_table.rows[0].cells[1].text = "Early Warning Signal"
            for cell in warn_table.rows[0].cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True

            # Data rows
            for i, warning in enumerate(unique_warnings):
                warn_table.rows[i + 1].cells[0].text = str(i + 1)
                warn_table.rows[i + 1].cells[1].text = str(warning)
        else:
            doc.add_paragraph("No significant early warning signals detected.", style="List Bullet")

        # Monitoring recommendations
        monitoring = safe_get(credit_decision, "monitoring_recommendations", [])
        if isinstance(monitoring, list) and monitoring:
            doc.add_paragraph("")
            add_sub_heading(doc, "Post-Disbursement Monitoring Recommendations")
            for rec in monitoring:
                doc.add_paragraph(str(rec), style="List Bullet")

        doc.add_paragraph("")

    except Exception as e:
        doc.add_paragraph(f"[Section 8 error: {str(e)}]")


# =============================================================
# SECTION 9: FINAL RECOMMENDATION
# =============================================================
def add_section_9_final_recommendation(doc, credit_decision, officer_assessment):
    """Section 9 — The final recommendation with bold decision banner."""
    try:
        add_section_heading(doc, "9. FINAL RECOMMENDATION")

        # Decision banner — big and bold
        decision = safe_get(credit_decision, "decision", "PENDING")
        overall_score = safe_get(credit_decision, "overall_score", "N/A")

        banner = doc.add_paragraph()
        banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = banner.add_run(f"\n{'━' * 50}\n")
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

        decision_para = doc.add_paragraph()
        decision_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = decision_para.add_run(f"FINAL RECOMMENDATION: {decision}")
        run.bold = True
        run.font.size = Pt(24)

        if "APPROVE" == str(decision).upper():
            run.font.color.rgb = RGBColor(0x27, 0xAE, 0x60)
        elif "REJECT" == str(decision).upper():
            run.font.color.rgb = RGBColor(0xE7, 0x4C, 0x3C)
        else:
            run.font.color.rgb = RGBColor(0xF3, 0x9C, 0x12)

        score_para = doc.add_paragraph()
        score_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = score_para.add_run(f"Overall Credit Score: {overall_score}/100")
        run.bold = True
        run.font.size = Pt(14)

        closing_banner = doc.add_paragraph()
        closing_banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = closing_banner.add_run(f"{'=' * 40}\n")
        run.font.size = Pt(12)

        # Score summary table — Five Cs at a glance
        add_sub_heading(doc, "Five Cs Score Summary")

        score_table = doc.add_table(rows=6, cols=3)
        score_table.style = "Light Grid Accent 1"

        # Header
        headers = ["Parameter", "Score", "Rating"]
        for j, header in enumerate(headers):
            score_table.rows[0].cells[j].text = header
            for paragraph in score_table.rows[0].cells[j].paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        # Five C scores
        c_names = [
            ("Character", "character_assessment"),
            ("Capacity", "capacity_assessment"),
            ("Capital", "capital_assessment"),
            ("Collateral", "collateral_assessment"),
            ("Conditions", "conditions_assessment")
        ]

        for i, (name, key) in enumerate(c_names):
            assessment = safe_get(credit_decision, key, {})
            score = safe_get(assessment, "score", "N/A") if isinstance(assessment, dict) else "N/A"
            rating = safe_get(assessment, "rating", "N/A") if isinstance(assessment, dict) else "N/A"

            score_table.rows[i + 1].cells[0].text = name
            score_table.rows[i + 1].cells[1].text = str(score)
            score_table.rows[i + 1].cells[2].text = str(rating)

        doc.add_paragraph("")

        # Recommended terms
        add_sub_heading(doc, "Recommended Loan Terms")
        doc.add_paragraph(f"Loan Amount: {safe_get(credit_decision, 'recommended_loan_amount', 'N/A')}", style="List Bullet")
        doc.add_paragraph(f"Interest Rate: {safe_get(credit_decision, 'recommended_interest_rate', 'N/A')}", style="List Bullet")
        doc.add_paragraph(f"Tenure: {safe_get(credit_decision, 'recommended_tenure', 'N/A')}", style="List Bullet")

        # Conditions for approval
        conditions = safe_get(credit_decision, "recommended_conditions", [])
        if isinstance(conditions, list) and conditions:
            add_sub_heading(doc, "Conditions for Approval")
            for condition in conditions:
                doc.add_paragraph(str(condition), style="List Bullet")

        # Top 3 reasons — THE MOST IMPORTANT PART
        add_sub_heading(doc, "Top Reasons for This Decision")
        reasons = safe_get(credit_decision, "top_reasons", [])
        if isinstance(reasons, list) and reasons:
            for i, reason in enumerate(reasons, 1):
                reason_para = doc.add_paragraph()
                run = reason_para.add_run(f"{i}. {reason}")
                run.bold = True
                run.font.size = Pt(11)
        else:
            doc.add_paragraph("No specific reasons available.", style="List Bullet")

        # Credit Officer assessment (if available)
        if officer_assessment and isinstance(officer_assessment, dict) and "error" not in officer_assessment:
            add_sub_heading(doc, "Credit Officer's Field Assessment")

            impression = safe_get(officer_assessment, "overall_officer_impression", {})
            if isinstance(impression, dict):
                doc.add_paragraph(f"Officer's Sentiment: {safe_get(impression, 'sentiment', 'N/A')}", style="List Bullet")
                doc.add_paragraph(f"Confidence in Borrower: {safe_get(impression, 'confidence_in_borrower', 'N/A')}", style="List Bullet")

            red_flags = safe_get(officer_assessment, "red_flags_from_visit", [])
            if isinstance(red_flags, list) and red_flags:
                doc.add_paragraph("Red Flags from Visit:")
                for flag in red_flags:
                    doc.add_paragraph(f"⚠️ {flag}", style="List Bullet 2")

        doc.add_paragraph("")

    except Exception as e:
        doc.add_paragraph(f"[Section 9 error: {str(e)}]")


# =============================================================
# FOOTER: Disclaimer
# =============================================================
def add_disclaimer(doc):
    """Adds the disclaimer and AI disclosure at the end."""
    try:
        doc.add_paragraph("━" * 60)

        disclaimer = doc.add_paragraph()
        run = disclaimer.add_run("DISCLAIMER")
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

        text = doc.add_paragraph()
        run = text.add_run(
            "This Credit Appraisal Memorandum has been generated using AI-assisted "
            "analysis (Intelli-Credit Engine). While every effort has been made to "
            "ensure accuracy, this report should be reviewed by a qualified credit "
            "officer before any lending decision is finalized. The AI recommendations "
            "are advisory in nature and do not constitute a binding credit decision. "
            "All data sources, including web research, are subject to the limitations "
            "of publicly available information."
        )
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        run.italic = True

        # Generated by line
        gen_para = doc.add_paragraph()
        gen_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        timestamp = datetime.now().strftime("%d %B %Y, %I:%M %p")
        run = gen_para.add_run(f"\nGenerated by Intelli-Credit AI Engine on {timestamp}")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
        run.italic = True

    except Exception as e:
        doc.add_paragraph(f"[Disclaimer error: {str(e)}]")


# =============================================================
# FORMATTING HELPERS
# =============================================================

def add_section_heading(doc, title):
    """Adds a styled section heading (e.g., '1. COMPANY OVERVIEW')."""
    heading = doc.add_heading(title, level=1)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)
        run.font.size = Pt(16)


def add_sub_heading(doc, title):
    """Adds a styled sub-heading within a section."""
    heading = doc.add_heading(title, level=2)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
        run.font.size = Pt(13)


def add_score_box(doc, label, assessment_data):
    """
    Adds a formatted score display for each of the Five Cs.

    Shows: Score number, Rating, Key Factors, and Explanation.
    """
    try:
        if not isinstance(assessment_data, dict):
            doc.add_paragraph(f"{label}: Data not available")
            return

        score = safe_get(assessment_data, "score", "N/A")
        rating = safe_get(assessment_data, "rating", "N/A")

        # Score and rating line
        score_para = doc.add_paragraph()
        run = score_para.add_run(f"{label}: {score}/100 ({rating})")
        run.bold = True
        run.font.size = Pt(13)

        # Color based on score
        try:
            score_num = int(score)
            if score_num >= 75:
                run.font.color.rgb = RGBColor(0x27, 0xAE, 0x60)
            elif score_num >= 60:
                run.font.color.rgb = RGBColor(0x29, 0x80, 0xB9)
            elif score_num >= 45:
                run.font.color.rgb = RGBColor(0xF3, 0x9C, 0x12)
            else:
                run.font.color.rgb = RGBColor(0xE7, 0x4C, 0x3C)
        except (ValueError, TypeError):
            pass

        # Key factors
        factors = safe_get(assessment_data, "key_factors", [])
        if isinstance(factors, list) and factors:
            doc.add_paragraph("Key Factors:")
            for factor in factors:
                doc.add_paragraph(str(factor), style="List Bullet")

        # Explanation
        explanation = safe_get(assessment_data, "explanation", "")
        if explanation:
            exp_para = doc.add_paragraph()
            run = exp_para.add_run(f"Analysis: {explanation}")
            run.italic = True
            run.font.size = Pt(10)

    except Exception as e:
        doc.add_paragraph(f"[Score box error for {label}: {str(e)}]")


def safe_get(data, key, default="N/A"):
    """
    Safely gets a value from a dictionary.
    Returns the default if the key doesn't exist or data isn't a dict.
    This prevents the report generation from crashing on missing data.

    Args:
        data:    The dictionary to look in
        key:     The key to find
        default: What to return if the key is missing

    Returns:
        The value if found, otherwise the default.
    """
    try:
        if isinstance(data, dict):
            return data.get(key, default)
        return default
    except Exception:
        return default
